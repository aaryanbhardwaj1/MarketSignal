"""Block renderers: native bar chart, compact table, metric callouts, status list and sources.

Each draws inside a ``Box`` (left, top, width, height in EMU). Charts show the engine's own
display strings as data labels; the emphasised bar uses the accent, the rest stay grey."""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any, cast

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.shapes.shapetree import SlideShapes
from pptx.util import Emu, Inches, Pt

from marketsignal.decks import theme as t
from marketsignal.decks.render_text import (
    cite_marker,
    claim_runs,
    line,
    line_break,
    new_paragraph,
    run,
    text_box,
)

ROW_H = Inches(0.32)
STATUS_ROW_H = Inches(0.56)
CAPTION_H = Inches(0.3)


@dataclass(frozen=True, slots=True)
class Box:
    left: int
    top: int
    width: int
    height: int


def solid(fill: Any, colour: RGBColor) -> None:
    fill.solid()
    fill.fore_color.rgb = colour


def hairline(shapes: SlideShapes, left: int, top: int, width: int) -> None:
    rule = shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Emu(left), Emu(top), Emu(left + width), Emu(top)
    )
    rule.line.color.rgb = t.HAIRLINE
    rule.line.width = Pt(0.75)


def caption(shapes: SlideShapes, box: Box, block: dict[str, Any]) -> None:
    parts = [p for p in (block.get("base"), block.get("note")) if p]
    text = " · ".join(parts) + cite_marker(block.get("cites", []))
    if text.strip():
        frame = text_box(shapes, Emu(box.left), Emu(box.top), Emu(box.width), CAPTION_H)
        line(frame, text.strip(), t.SMALL_PT, t.MUTED)


def chart(shapes: SlideShapes, box: Box, block: dict[str, Any]) -> None:
    data: Any = cast(Any, CategoryChartData)()  # unannotated in python-pptx
    data.categories = block["categories"]
    data.add_series(block["title"], block["values"])
    height = box.height - CAPTION_H
    frame = shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Emu(box.left),
        Emu(box.top),
        Emu(box.width),
        Emu(height),
        data,
    )
    ch = cast(Any, frame).chart
    ch.has_legend = False
    ch.font.name = t.FONT
    ch.font.size = t.SMALL_PT
    ch.font.color.rgb = t.INK
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.vary_by_categories = False
    category = ch.category_axis
    category.reverse_order = True
    category.has_major_gridlines = False
    category.format.line.color.rgb = t.HAIRLINE
    value = ch.value_axis
    value.visible = False
    value.has_major_gridlines = False
    low = min(block["values"])
    value.minimum_scale = min(0.0, low * 1.15)
    value.maximum_scale = max(block["values"]) * 1.18 or 1.0
    for i, point in enumerate(plot.series[0].points):
        solid(point.format.fill, t.ACCENT if i == block["emphasis"] else t.BAR_GREY)
        label = point.data_label
        label.position = XL_LABEL_POSITION.OUTSIDE_END
        label.text_frame.text = block["displays"][i]
    caption(shapes, Box(box.left, box.top + height, box.width, CAPTION_H), block)


def table(shapes: SlideShapes, box: Box, block: dict[str, Any]) -> None:
    n_rows = len(block["rows"]) + 1
    n_cols = len(block["columns"])
    height = min(box.height - CAPTION_H, ROW_H * n_rows)
    frame = shapes.add_table(
        n_rows, n_cols, Emu(box.left), Emu(box.top), Emu(box.width), Emu(height)
    )
    tbl = frame.table
    tbl.first_row = True
    tbl.horz_banding = False
    cells = [block["columns"], *block["rows"]]
    for r, values in enumerate(cells):
        tbl.rows[r].height = Emu(height // n_rows)
        for c, value in enumerate(values):
            cell = tbl.cell(r, c)
            solid(cell.fill, t.SUNKEN if r == 0 else t.WHITE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = cell.text_frame.paragraphs[0]
            numeric = c == n_cols - 1 and r > 0 and block["columns"][-1] == "Value"
            paragraph.alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
            run(paragraph, str(value), t.SMALL_PT, t.INK, bold=r == 0)
    caption(shapes, Box(box.left, box.top + height + Inches(0.08), box.width, CAPTION_H), block)


def metrics(shapes: SlideShapes, box: Box, block: dict[str, Any]) -> None:
    cols = 3
    cell_w = box.width // cols
    cell_h = Inches(2.1)
    for i, item in enumerate(block["items"]):
        left = box.left + (i % cols) * cell_w
        top = box.top + (i // cols) * (cell_h + Inches(0.3))
        hairline(shapes, left, top, cell_w - Inches(0.3))
        frame = text_box(
            shapes, Emu(left), Emu(top + Inches(0.15)), Emu(cell_w - Inches(0.3)), Emu(cell_h)
        )
        line(frame, item["display"], t.METRIC_PT, t.INK, bold=True)
        line(frame, item["label"], t.SUB_BULLET_PT, t.INK, first=False)
        base = (item.get("base") or "") + cite_marker(item.get("cites", []))
        if base.strip():
            line(frame, base.strip(), t.SMALL_PT, t.MUTED, first=False)


def statuses(shapes: SlideShapes, box: Box, block: dict[str, Any]) -> None:
    status_w = Inches(2.4)
    for i, item in enumerate(block["items"]):
        top = box.top + i * STATUS_ROW_H
        title_w = box.width - status_w
        frame = text_box(
            shapes, Emu(box.left), Emu(top), Emu(title_w), STATUS_ROW_H, anchor=MSO_ANCHOR.MIDDLE
        )
        paragraph = frame.paragraphs[0]
        claim_runs(paragraph, {**item, "text": item["title"]}, t.BODY_PT)
        meta = text_box(
            shapes,
            Emu(box.left + title_w),
            Emu(top),
            status_w,
            STATUS_ROW_H,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        label = item.get("status_label")
        colour = t.STATUS_COLOURS.get(label or "", t.MUTED)
        line(meta, label or item.get("kind_label") or "", t.SUB_BULLET_PT, colour, bold=bool(label))
        meta.paragraphs[0].alignment = PP_ALIGN.RIGHT
        hairline(shapes, box.left, top + STATUS_ROW_H, box.width)


def sources(shapes: SlideShapes, box: Box, block: dict[str, Any]) -> None:
    frame = text_box(
        shapes, Emu(box.left), Emu(box.top), Emu(box.width), Emu(box.height), shrink=True
    )
    if not block["items"]:
        line(frame, "No sources were cited in this deck.", t.SUB_BULLET_PT, t.MUTED)
        return
    for i, item in enumerate(block["items"]):
        paragraph = frame.paragraphs[0] if i == 0 else new_paragraph(frame)
        paragraph.space_before = Pt(8 if i else 0)
        run(paragraph, f"[{item['n']}] ", t.SMALL_PT, t.ACCENT, bold=True)
        where = ", ".join(p for p in (item["source_title"], item.get("locator")) if p)
        run(paragraph, where, t.SMALL_PT, t.INK, bold=True)
        if item.get("quote"):
            line_break(paragraph)
            run(paragraph, f"“{item['quote']}”", t.SMALL_PT, t.MUTED, italic=True)
