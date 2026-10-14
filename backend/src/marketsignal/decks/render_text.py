"""Text primitives for the deck renderer: text boxes, runs, bullet paragraphs with claim tags
and citation markers. Raw handles never reach these helpers; citations are numbers only."""

from __future__ import annotations

from typing import Any, cast

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.shapes.shapetree import SlideShapes
from pptx.text.text import TextFrame, _Paragraph
from pptx.util import Emu, Length, Pt

from marketsignal.decks import theme as t


def text_box(
    shapes: SlideShapes,
    left: Length,
    top: Length,
    width: Length,
    height: Length,
    *,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    shrink: bool = False,
) -> TextFrame:
    frame: TextFrame = shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Emu(0)
    frame.margin_top = frame.margin_bottom = Emu(0)
    frame.vertical_anchor = anchor
    if shrink:
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return frame


def new_paragraph(frame: TextFrame) -> _Paragraph:
    paragraph: _Paragraph = cast(Any, frame).add_paragraph()  # unannotated in python-pptx
    return paragraph


def line_break(paragraph: _Paragraph) -> None:
    cast(Any, paragraph).add_line_break()  # unannotated in python-pptx


def run(
    paragraph: _Paragraph,
    text: str,
    size: Length,
    colour: RGBColor = t.INK,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    r = paragraph.add_run()
    r.text = text
    font = r.font
    font.name = t.FONT
    font.size = size
    font.bold = bold
    font.italic = italic
    font.color.rgb = colour


def line(
    frame: TextFrame,
    text: str,
    size: Length,
    colour: RGBColor = t.INK,
    *,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    first: bool = True,
) -> _Paragraph:
    paragraph = frame.paragraphs[0] if first else new_paragraph(frame)
    paragraph.alignment = align
    run(paragraph, text, size, colour, bold=bold)
    return paragraph


def cite_marker(cites: list[int]) -> str:
    return f" [{', '.join(str(n) for n in cites)}]" if cites else ""


def _bullet_glyph(paragraph: _Paragraph, level: int) -> None:
    indent = Emu(Pt(14) + level * Pt(22))
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(int(indent)))
    p_pr.set("indent", str(-int(Pt(14))))
    colour = p_pr.makeelement(qn("a:buClr"), {})
    colour.append(colour.makeelement(qn("a:srgbClr"), {"val": str(t.ACCENT)}))
    p_pr.append(colour)
    p_pr.append(p_pr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "\u2013"}))


def claim_runs(paragraph: _Paragraph, item: dict[str, Any], size: Length) -> None:
    """Optional claim tag, the text, then the citation numbers."""
    label = item.get("label")
    if label:
        run(paragraph, f"{label}: ", size, t.CLAIM_COLOURS.get(label, t.MUTED), bold=True)
    run(paragraph, item["text"], size, t.INK if item.get("level", 0) == 0 else t.MUTED)
    marker = cite_marker(item.get("cites", []))
    if marker:
        run(paragraph, marker, t.SMALL_PT, t.SUBTLE)


def bullets(frame: TextFrame, items: list[dict[str, Any]]) -> None:
    for i, item in enumerate(items):
        level = int(item.get("level", 0))
        paragraph = frame.paragraphs[0] if i == 0 else new_paragraph(frame)
        paragraph.space_before = Pt(10 if level == 0 and i else 3)
        paragraph.space_after = Pt(2)
        size = t.BODY_PT if level == 0 else t.SUB_BULLET_PT
        if item.get("heading"):
            run(paragraph, item["heading"], size, t.INK, bold=True)
            line_break(paragraph)
        claim_runs(paragraph, item, size)
        _bullet_glyph(paragraph, level)
