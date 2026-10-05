"""CSV and XLSX parser: rows are citation units, typed tables feed analytics.

* The header is the first non-empty row; header cells must be unique and non-empty.
* Row numbers are spreadsheet numbering (header = row 1 for CSV), so a cited row can be found
  by opening the file: CSV ``R{n}``, XLSX ``SH{sheet}.R{n}``.
* Each row's parent text is the authoritative ``column: value; ...`` rendering.
* Column roles (plan §9): ``free_text`` (long or high-cardinality strings), ``identifier``,
  ``context`` (low-cardinality strings), ``constant`` (one repeated value), ``numeric``, ``date``.
* Only rows with a non-empty free-text cell become retrieval units. Their child text is
  ``{table} | context=values | column: verbatim``, and their span points at the verbatim value
  inside the parent text (D1). Purely numeric rows are parents only: resolvable, available to
  analytics, but not retrieval units (they would flood both lanes with near-identical text).
* One table-summary parent per table (``T1`` / ``SH{n}.T1``) describes the columns and levels.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from collections import Counter
from collections.abc import Iterable
from typing import Any

from openpyxl import load_workbook

from marketsignal.config import Settings
from marketsignal.domain.enums import IngestErrorCode
from marketsignal.evidence.handles import LocatorKind
from marketsignal.ingestion.models import (
    ColumnProfile,
    DatasetTable,
    IngestionError,
    ParentDraft,
    ParentKind,
    ParsedSource,
    RowChildSpec,
)
from marketsignal.ingestion.structure import clean
from marketsignal.ingestion.tokenizer import Tokenizer
from marketsignal.ingestion.validation import decode_text

_NUMBER = re.compile(r"^[-+]?(\d+(\.\d*)?|\.\d+)([eE][-+]?\d+)?$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}(-\d{2})?$")
FREE_TEXT_MEAN_LEN = 20
IDENTIFIER_MAX_LEN = 64
FREE_TEXT_MULTIWORD_SHARE = 0.5  # free text is prose: most values have several words
NUMERIC_SHARE = 0.9


def _malformed(message: str) -> IngestionError:
    return IngestionError(IngestErrorCode.MALFORMED_SPREADSHEET.value, message)


def _normalise(value: Any) -> Any:
    """Cell value as stored in dataset_rows (JSON-safe) and rendered in parent text."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, int | float):
        return value
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = clean(str(value))
    return text or None


def _coerce_number(value: Any) -> Any:
    if isinstance(value, str) and _NUMBER.match(value):
        number = float(value)
        return int(number) if number.is_integer() and "." not in value else number
    return value


def _profile(name: str, values: list[Any], max_levels: int) -> ColumnProfile:
    present = [v for v in values if v is not None]
    distinct = len(set(map(str, present)))
    if not present:
        return ColumnProfile(name, "text", "context", 0, 0)
    numeric = sum(isinstance(v, int | float) and not isinstance(v, bool) for v in present)
    if numeric / len(present) >= NUMERIC_SHARE:
        return ColumnProfile(name, "numeric", "numeric", len(present), distinct)
    if all(isinstance(v, str) and _ISO_DATE.match(v) for v in present):
        return ColumnProfile(name, "date", "date", len(present), distinct)
    if distinct == 1 and len(present) > 1:
        # A single repeated value (e.g. a data notice) carries no row-specific evidence.
        return ColumnProfile(
            name, "categorical", "constant", len(present), distinct, (str(present[0]),)
        )
    mean_len = sum(len(str(v)) for v in present) / len(present)
    tokens_only = all(not any(ch.isspace() for ch in str(v)) for v in present)
    # Identifiers first: unique, whitespace-free tokens of any length (e.g. CP-2025-10-SITE-GENZ)
    # must never be mistaken for free text, or every numeric row would become a retrieval unit.
    if (
        distinct == len(present)
        and len(present) > 1
        and tokens_only
        and mean_len <= IDENTIFIER_MAX_LEN
    ):
        return ColumnProfile(name, "text", "identifier", len(present), distinct)
    multi_word = (
        sum(" " in str(v).strip() for v in present) / len(present) >= FREE_TEXT_MULTIWORD_SHARE
    )
    if mean_len > FREE_TEXT_MEAN_LEN and multi_word:
        return ColumnProfile(name, "text", "free_text", len(present), distinct)
    levels = tuple(level for level, _ in Counter(map(str, present)).most_common(max_levels))
    return ColumnProfile(name, "categorical", "context", len(present), distinct, levels)


def _render_value(value: Any) -> str:
    return str(value)


def _build_table(
    *,
    rows: Iterable[tuple[int, list[Any]]],
    sheet_ordinal: int,
    sheet_name: str,
    title: str,
    locator_prefix: tuple[tuple[LocatorKind, int], ...],
    settings: Settings,
) -> tuple[DatasetTable, list[ParentDraft]] | None:
    # (row number, original cells for rendering, typed cells for profiling/analytics)
    materialised: list[tuple[int, list[Any], list[Any]]] = []
    header: list[str] | None = None
    header_row = 0
    for number, cells in rows:
        cells = [_normalise(c) for c in cells]
        if header is None:
            if not any(c is not None for c in cells):
                continue
            while cells and cells[-1] is None:
                cells.pop()
            if any(c is None for c in cells):
                raise _malformed(f"empty header cell in {sheet_name!r}")
            header = [str(c) for c in cells]
            if len(set(header)) != len(header):
                raise _malformed(f"duplicate header names in {sheet_name!r}")
            header_row = number
            continue
        if not any(c is not None for c in cells):
            continue
        if len(cells) > len(header) and any(c is not None for c in cells[len(header) :]):
            raise _malformed(f"row {number} of {sheet_name!r} has more cells than the header")
        raw = cells[: len(header)]
        materialised.append((number, raw, [_coerce_number(c) for c in raw]))
        if len(materialised) > settings.max_table_rows:
            raise IngestionError(IngestErrorCode.CONTENT_TOO_LARGE.value, "too many rows")
    if header is None:
        return None

    columns = tuple(
        _profile(
            name,
            [row[i] if i < len(row) else None for _, _raw, row in materialised],
            settings.table_summary_max_levels,
        )
        for i, name in enumerate(header)
    )
    roles = {c.name: c.role for c in columns}
    table_rows: list[tuple[int, dict[str, Any]]] = []
    parents: list[ParentDraft] = []
    for number, raw, row in materialised:
        values = {name: (row[i] if i < len(row) else None) for i, name in enumerate(header)}
        originals = {name: (raw[i] if i < len(raw) else None) for i, name in enumerate(header)}
        table_rows.append((number, values))
        parts: list[str] = []
        free_span: tuple[int, int, str, str] | None = None
        offset = 0
        for name in header:
            value = originals[name]
            if value is None:
                continue
            prefix = f"{name}: "
            rendered = _render_value(value)  # authoritative original text, never coerced
            start = offset + len(prefix)
            if roles[name] == "free_text" and free_span is None:
                free_span = (start, start + len(rendered), name, rendered)
            parts.append(prefix + rendered)
            offset += len(prefix) + len(rendered) + 2  # "; " separator
        text = "; ".join(parts)
        row_child = None
        if free_span is not None:
            start, end, column, verbatim = free_span
            identifiers = (
                [f"{n}={originals[n]}" for n in header if roles[n] == "identifier" and originals[n]]
                if settings.row_child_identifiers
                else []
            )
            context = ", ".join(
                identifiers
                + [
                    f"{n}={values[n]}"
                    for n in header
                    if roles[n] == "context" and values[n] is not None
                ]
            )
            retrieval = " | ".join(p for p in (f"{title} > {sheet_name}", context) if p)
            row_child = RowChildSpec(f"{retrieval} | {column}: {verbatim}", start, end)
        parents.append(
            ParentDraft(
                locator=(*locator_prefix, (LocatorKind.ROW, number)),
                kind=ParentKind.ROW,
                text=text,
                heading_path=(sheet_name,),
                locator_meta={"sheet": sheet_ordinal, "sheet_name": sheet_name, "row": number},
                row_child=row_child,
                index_children=row_child is not None,
            )
        )
    summary = _summary_text(sheet_name, title, columns, len(materialised))
    parents.insert(
        0,
        ParentDraft(
            locator=(*locator_prefix, (LocatorKind.TABLE, 1)),
            kind=ParentKind.TABLE_SUMMARY,
            text=summary,
            heading_path=(sheet_name,),
            locator_meta={"sheet": sheet_ordinal, "sheet_name": sheet_name, "table": 1},
        ),
    )
    table = DatasetTable(sheet_ordinal, sheet_name, header_row, columns, tuple(table_rows))
    return table, parents


def _summary_text(
    sheet_name: str, title: str, columns: tuple[ColumnProfile, ...], row_count: int
) -> str:
    described = []
    for c in columns:
        if c.role == "context" and c.levels:
            described.append(f"{c.name} (categorical: {', '.join(c.levels)})")
        elif c.role == "free_text":
            described.append(f"{c.name} (free text)")
        else:
            described.append(f"{c.name} ({c.role})")
    return (
        f"Table '{sheet_name}' in {title}: {row_count} data rows. Columns: {'; '.join(described)}."
    )


def parse_csv(data: bytes, tokenizer: Tokenizer, settings: Settings, *, title: str) -> ParsedSource:
    del tokenizer  # rows are atomic citation units; no token-based grouping
    text = decode_text(data)
    try:
        reader = list(csv.reader(io.StringIO(text)))
    except csv.Error as exc:
        raise _malformed(f"unreadable CSV: {exc}") from exc
    built = _build_table(
        rows=((i, list(r)) for i, r in enumerate(reader, start=1)),
        sheet_ordinal=1,
        sheet_name=title,
        title=title,
        locator_prefix=(),
        settings=settings,
    )
    if built is None:
        return ParsedSource(parents=())
    table, parents = built
    return ParsedSource(parents=tuple(parents), tables=(table,))


def parse_xlsx(
    data: bytes, tokenizer: Tokenizer, settings: Settings, *, title: str
) -> ParsedSource:
    del tokenizer
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types on corrupt input
        raise IngestionError(IngestErrorCode.PARSE_FAILED.value, "unreadable XLSX") from exc
    parents: list[ParentDraft] = []
    tables: list[DatasetTable] = []
    try:
        for ordinal, sheet in enumerate(workbook.worksheets, start=1):
            built = _build_table(
                rows=(
                    (i, list(r)) for i, r in enumerate(sheet.iter_rows(values_only=True), start=1)
                ),
                sheet_ordinal=ordinal,
                sheet_name=sheet.title,
                title=title,
                locator_prefix=((LocatorKind.SHEET, ordinal),),
                settings=settings,
            )
            if built is not None:
                tables.append(built[0])
                parents.extend(built[1])
    finally:
        workbook.close()
    return ParsedSource(parents=tuple(parents), tables=tuple(tables))
