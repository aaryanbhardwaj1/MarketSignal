"""PPTX parser (python-pptx).

Each slide is one parent ``SL{n}`` (title + text frames + tables, in shape order, including
grouped shapes); speaker notes are a separate parent ``SL{n}.N1``. A slide whose text exceeds
the parent cap is split into ``SL{n}.B{k}`` blocks (it never merges with another slide).
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from marketsignal.config import Settings
from marketsignal.domain.enums import IngestErrorCode
from marketsignal.evidence.handles import LocatorKind
from marketsignal.ingestion.models import IngestionError, ParentDraft, ParentKind, ParsedSource
from marketsignal.ingestion.structure import Unit, clean, group_units
from marketsignal.ingestion.tokenizer import Tokenizer


def _shapes(shapes: Any) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _shape_lines(shape: Any) -> list[str]:
    lines: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = clean("".join(run.text for run in paragraph.runs) or paragraph.text)
            if text:
                lines.append(("- " if paragraph.level > 0 else "") + text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [clean(cell.text) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return lines


def parse_pptx(data: bytes, tokenizer: Tokenizer, settings: Settings) -> ParsedSource:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # python-pptx raises several unrelated types on corrupt input
        raise IngestionError(IngestErrorCode.PARSE_FAILED.value, "unreadable PPTX") from exc
    slides = list(presentation.slides)
    if len(slides) > settings.max_slides:
        raise IngestionError(IngestErrorCode.CONTENT_TOO_LARGE.value, "too many slides")
    parents: list[ParentDraft] = []
    for number, slide in enumerate(slides, start=1):
        title_shape = slide.shapes.title
        title = clean(title_shape.text) if title_shape is not None else ""
        body: list[str] = []
        for shape in _shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            body.extend(_shape_lines(shape))
        heading = (title,) if title else ()
        meta = {"slide": number, "slide_title": title}
        slide_text = "\n".join(([title] if title else []) + body)
        if slide_text:
            parents.extend(
                _split_container(
                    ((LocatorKind.SLIDE, number),),
                    ParentKind.SLIDE,
                    slide_text,
                    heading,
                    meta,
                    tokenizer,
                    settings.parent_max_tokens,
                )
            )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            notes_text = "\n".join(
                clean(p.text)
                for p in (notes.paragraphs if notes is not None else [])
                if clean(p.text)
            )
            if notes_text:
                parents.extend(
                    _split_container(
                        ((LocatorKind.SLIDE, number), (LocatorKind.NOTES, 1)),
                        ParentKind.NOTES,
                        notes_text,
                        heading,
                        {**meta, "notes": True},
                        tokenizer,
                        settings.parent_max_tokens,
                    )
                )
    return ParsedSource(parents=tuple(parents))


def _split_container(
    locator: tuple[tuple[LocatorKind, int], ...],
    kind: ParentKind,
    text: str,
    heading: tuple[str, ...],
    meta: dict[str, Any],
    tokenizer: Tokenizer,
    max_tokens: int,
) -> list[ParentDraft]:
    if tokenizer.count(text) <= max_tokens:
        return [ParentDraft(locator, kind, text, heading, meta)]
    blocks = group_units([Unit(line) for line in text.split("\n")], tokenizer, max_tokens, "")
    return [
        ParentDraft((*locator, (LocatorKind.BLOCK, i)), kind, block.text, heading, meta)
        for i, block in enumerate(blocks, start=1)
    ]
