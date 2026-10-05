"""Build small documents of every supported format at test time (no committed binaries)."""

from __future__ import annotations

import csv
import io
import zipfile

import docx
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen.canvas import Canvas

BODY = 10.5
LINE = 13.0


def make_pdf(
    pages: list[list[tuple[str, str]]], *, footer: str | None = "Fictional - page"
) -> bytes:
    """pages: list of [(kind, text)] where kind in {"h1", "h2", "p", "bullet", "gap"}."""
    buffer = io.BytesIO()
    canvas = Canvas(buffer, pagesize=LETTER, invariant=1)
    for number, items in enumerate(pages, start=1):
        y = 740.0
        for kind, text in items:
            if kind == "gap":
                y -= 10
                continue
            font, size = {
                "h1": ("Helvetica-Bold", 18),
                "h2": ("Helvetica-Bold", 14),
                "p": ("Helvetica", BODY),
                "bullet": ("Helvetica", BODY),
            }[kind]
            canvas.setFont(font, size)
            for line in text.split("\n"):
                canvas.drawString(72, y, ("- " + line) if kind == "bullet" else line)
                y -= size + 3
            if kind in ("h1", "h2"):
                y -= 4
        if footer:
            canvas.setFont("Helvetica", 7)
            canvas.drawString(72, 40, f"{footer} {number}")
        canvas.showPage()
    canvas.save()
    return buffer.getvalue()


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    """blocks: (style, text) where style in {"h1", "h2", "p", "bullet", "table:a|b;c|d"}."""
    document = docx.Document()
    for style, text in blocks:
        if style == "h1":
            document.add_heading(text, level=1)
        elif style == "h2":
            document.add_heading(text, level=2)
        elif style == "bullet":
            document.add_paragraph(text, style="List Bullet")
        elif style.startswith("table"):
            rows = [r.split("|") for r in text.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx(slides: list[dict[str, object]]) -> bytes:
    """slides: {"title": str, "bullets": [str], "table": [[str]] | None, "notes": str | None}."""
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # title + content
    for spec in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = str(spec["title"])
        body = slide.placeholders[1].text_frame
        bullets = list(spec.get("bullets") or [])  # type: ignore[call-overload]
        if bullets:
            body.text = bullets[0]
            for bullet in bullets[1:]:
                body.add_paragraph().text = bullet
        table = spec.get("table")
        if table:
            rows = list(table)  # type: ignore[call-overload]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(1), Inches(4.5), Inches(6), Inches(1)
            )
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
        notes = spec.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = str(notes)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_csv(rows: list[list[object]], *, bom: bool = False) -> bytes:
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerows(rows)
    return (("﻿" if bom else "") + buffer.getvalue()).encode("utf-8")


def make_zip(entries: dict[str, bytes], *, compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=compression) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


CONTENT_TYPES_DOCX = (
    b'<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
    b'content-types"><Override PartName="/word/document.xml" ContentType="application/'
    b'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
)
CONTENT_TYPES_MACRO = (
    b'<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
    b'content-types"><Override PartName="/word/document.xml" ContentType="application/'
    b'vnd.ms-word.document.macroEnabled.main+xml"/></Types>'
)
