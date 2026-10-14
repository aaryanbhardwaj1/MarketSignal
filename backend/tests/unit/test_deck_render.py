"""PowerPoint rendering of a composed deck (in memory) and the safe download name."""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Inches

from marketsignal.decks import render
from marketsignal.decks.filename import content_disposition, safe_filename
from marketsignal.decks.render import DeckTooLargeError, render_pptx
from tests.unit.test_deck_composer import analysis, build, metric, opportunity, ref, signal, world


def _deck() -> dict[str, Any]:
    outlook = signal(
        "s2",
        "future_outlook",
        claim_type="forecast",
        evidence=[ref("ACME/CALL@v1:S1", "Below 5% through 2028.", title="Analyst call")],
    )
    brief, listing = world(
        [
            signal(
                "s1",
                metrics=[metric("a-1", "Share named fit", 27.0)],
                implication={"text": "Fix fit.", "claim_type": "inference"},
            ),
            outlook,
        ],
        [opportunity("O1", contradicting=[ref("ACME/C@v1:P1", "C.", "contradict")])],
        [analysis("a-1"), analysis("a-2", 12, presentation="table")],
    )
    return build(
        brief, listing, signal_ids=["s1", "s2"], opportunity_ids=["O1"], analysis_ids=["a-1", "a-2"]
    )


def _texts(slide: Any) -> list[str]:
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out += [cell.text for row in shape.table.rows for cell in row.cells]
    return out


def test_renders_every_slide_with_titles_notes_chart_and_appendix() -> None:
    deck = _deck()
    prs = Presentation(BytesIO(render_pptx(deck)))
    assert (prs.slide_width, prs.slide_height) == (Inches(13.333), Inches(7.5))
    slides = list(prs.slides)
    assert len(slides) == len(deck["slides"])
    for spec, slide in zip(deck["slides"], slides, strict=True):
        texts = _texts(slide)
        assert spec["title"] in texts, spec["title"]
        joined = "\n".join(texts)
        assert "ACME/" not in joined
        assert "@v1" not in joined
        if spec["notes"]:
            notes = slide.notes_slide.notes_text_frame.text
            for note in spec["notes"]:
                assert f"[{note['n']}]" in notes
                assert note["handle"] in notes
    by_title = {s["title"]: slides[s["index"] - 1] for s in deck["slides"]}
    charts = [sh for sh in by_title["Analysis a-1"].shapes if sh.has_chart]
    assert len(charts) == 1
    plot = charts[0].chart.plots[0]
    assert list(plot.categories) == ["Pain 0", "Pain 1", "Pain 2"]
    tables = [sh for sh in by_title["Analysis a-2"].shapes if sh.has_table]
    assert len(tables) == 1
    assert len(tables[0].table.rows) == 11  # header + ten rows shown
    appendix = "\n".join(_texts(by_title["Appendix: Sources"]))
    assert "Analyst call" in appendix
    assert "Below 5% through 2028." in appendix
    market = "\n".join(_texts(by_title["Competitor and market findings"]))
    assert "Forecast:" in market
    customer = "\n".join(_texts(by_title["Key customer findings"]))
    assert "Inference:" in customer
    assert "MarketSignal" in customer


def test_render_is_deterministic() -> None:
    deck = _deck()
    assert render_pptx(deck) == render_pptx(deck)


def test_oversized_output_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(render, "MAX_DECK_BYTES", 1000)
    with pytest.raises(DeckTooLargeError):
        render_pptx(_deck())


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("Strategic findings: Northstar", "Strategic-findings-Northstar.pptx"),
        ('a"b\r\nc/../d\\e', "a-b-c-d-e.pptx"),
        ("Café — growth", "Cafe-growth.pptx"),
        ("../../etc/passwd", "etc-passwd.pptx"),
        ("", "marketsignal-deck.pptx"),
        (None, "marketsignal-deck.pptx"),
        ("日本語", "marketsignal-deck.pptx"),
        ("x" * 200, "x" * 80 + ".pptx"),
    ],
)
def test_safe_filename(title: str | None, expected: str) -> None:
    assert safe_filename(title) == expected


def test_content_disposition_is_quoted_ascii() -> None:
    header = content_disposition('Q3 "deck"\n')
    assert header == 'attachment; filename="Q3-deck.pptx"'
    assert header.isascii()
