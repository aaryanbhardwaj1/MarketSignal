"""Deck preview and PowerPoint download over harness-uploaded sources, with the synthetic v2
fixture injected through the insights loader: grounded composition from the re-verified brief,
a real .pptx (titles, notes, native chart, appendix), validation, unknown and foreign ids, and
items that stop verifying after the corpus changes."""

# ruff: noqa: F811
from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest
from pptx import Presentation

from tests.integration.test_ingestion_api import Harness, harness  # noqa: F401
from tests.integration.test_insights_api import FIT, GONE, _source_id, _upload
from tests.integration.test_insights_v2_api import PROJ, RISK, world  # noqa: F401

pytestmark = pytest.mark.integration

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
FULL = {
    "signal_ids": ["fit", "outlook"],
    "opportunity_ids": ["O1"],
    "analysis_ids": ["score-by-segment", "score-drifted"],
}


async def _preview(h: Harness, ws: str, body: dict[str, Any]) -> Any:
    return await h.client.post(f"/api/workspaces/{ws}/decks/preview", json=body)


def _texts(slide: Any) -> str:
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out += [c.text for row in shape.table.rows for c in row.cells]
    return "\n".join(t for t in out if t)


async def test_preview_composes_from_the_verified_brief(world: tuple[Harness, str]) -> None:
    h, ws = world
    response = await _preview(h, ws, FULL)
    assert response.status_code == 200, response.text
    deck = response.json()["deck"]
    assert deck["title"] == f"Strategic findings: {ws}"
    assert deck["workspace"] == {"code": ws, "name": ws}
    assert deck["fixture_version"] == f"{ws}.v2test"
    assert deck["excluded"] == [
        {
            "id": "score-drifted",
            "reason": "Its figures could not be recomputed from the current data.",
        }
    ]
    kinds = [s["kind"] for s in deck["slides"]]
    assert kinds == [
        "title",
        "summary",
        "customer_findings",
        "quantitative",
        "quantitative",
        "market_findings",
        "opportunity",
        "risks",
        "next_steps",
        "appendix",
    ]
    assert deck["counts"] == {"items": 4, "sources": len(deck["sources"]), "slides": 10}
    quotes = {s["quote"] for s in deck["sources"] if s["kind"] == "passage"}
    assert quotes == {FIT, PROJ, GONE, RISK}
    market = next(s for s in deck["slides"] if s["kind"] == "market_findings")
    labels = {b["text"]: b["label"] for b in market["blocks"][0]["items"]}
    assert labels["Personalised units stay a small share of volume."] == "Forecast"
    assert labels["Do not plan around it."] == "Hypothesis"
    chart_slide = next(s for s in deck["slides"] if s["title"] == "Score by segment")
    assert chart_slide["blocks"][0]["type"] == "chart"
    assert chart_slide["blocks"][0]["displays"] == ["2.00", "5.00"]


async def test_pptx_download_is_a_real_presentation(world: tuple[Harness, str]) -> None:
    h, ws = world
    body = {**FULL, "title": 'Fit "deck"\r\n/../x'}
    response = await h.client.post(f"/api/workspaces/{ws}/decks/pptx", json=body)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == PPTX
    assert response.headers["content-disposition"] == 'attachment; filename="Fit-deck-x.pptx"'
    prs = Presentation(BytesIO(response.content))
    slides = list(prs.slides)
    titles = [_texts(s).split("\n")[0] for s in slides]
    assert titles[0] == 'Fit "deck" /../x'
    assert "Executive summary" in titles
    assert "Appendix: Sources" in titles
    chart_slide = slides[titles.index("Score by segment")]
    assert any(shape.has_chart for shape in chart_slide.shapes)
    customer = slides[titles.index("Key customer findings")]
    notes = customer.notes_slide.notes_text_frame.text
    assert FIT in notes
    assert "(handle: " in notes
    assert "@v" not in _texts(customer)
    assert FIT in _texts(slides[titles.index("Appendix: Sources")])


async def test_validation_and_unknown_ids(world: tuple[Harness, str]) -> None:
    h, ws = world
    empty = await _preview(h, ws, {"signal_ids": []})
    assert empty.status_code == 422
    assert empty.json()["error"]["code"] == "VALIDATION_ERROR"
    extra = await _preview(h, ws, {"signal_ids": ["fit"], "workspace": "OTHER"})
    assert extra.status_code == 422
    unknown = await _preview(h, ws, {"signal_ids": ["fit", "nope"], "analysis_ids": ["zzz"]})
    assert unknown.status_code == 422
    error = unknown.json()["error"]
    assert error["code"] == "DECK_ITEMS_NOT_FOUND"
    assert error["unknown"] == ["nope", "zzz"]


async def test_ids_from_another_workspace_are_not_usable(world: tuple[Harness, str]) -> None:
    h, _ = world
    other = await h.create_workspace()
    response = await _preview(h, other, {"signal_ids": ["fit"], "opportunity_ids": ["O1"]})
    assert response.status_code == 422
    assert response.json()["error"]["unknown"] == ["fit", "O1"]
    download = await h.client.post(f"/api/workspaces/{other}/decks/pptx", json=FULL)
    assert download.status_code == 422


async def test_items_that_stop_verifying_are_excluded(world: tuple[Harness, str]) -> None:
    h, ws = world
    await _upload(h, ws, "risks.md", b"# Risks\n\nThe risk register was rewritten.\n", "RISKS")
    purge = await h.client.delete(f"/api/workspaces/{ws}/sources/{await _source_id(h, ws, 'GONE')}")
    assert purge.status_code == 200, purge.text
    deck = (await _preview(h, ws, FULL)).json()["deck"]
    quotes = {s["quote"] for s in deck["sources"]}
    assert GONE not in quotes
    assert RISK not in quotes
    assert "risks" not in [s["kind"] for s in deck["slides"]]
    opp = next(s for s in deck["slides"] if s["kind"] == "opportunity")
    texts = [b["text"] for b in opp["blocks"][0]["items"]]
    assert "1 supporting and 0 contradicting passage" in texts

    await _upload(h, ws, "memo.md", b"# Notes\n\nEverything was rewritten.\n", "MEMO")
    gone = await _preview(h, ws, {"signal_ids": ["fit"]})
    assert gone.status_code == 200
    assert [e["id"] for e in gone.json()["deck"]["excluded"]] == ["fit"]
    empty = await h.client.post(f"/api/workspaces/{ws}/decks/pptx", json={"signal_ids": ["fit"]})
    assert empty.status_code == 422
    assert empty.json()["error"]["code"] == "DECK_EMPTY"
