"""Deterministic renderers for narrative / slide / tabular documents.

Fictional data created for the MarketSignal demo.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from common import DISCLAIMER, DISCLAIMER_CORE

FIXED_DT = dt.datetime(2026, 1, 1, 0, 0, 0)
ZIP_DATE = (2026, 1, 1, 0, 0, 0)
AUTHOR = "MarketSignal seed generator"


# ---------------------------------------------------------------------------
# Zip normalisation (DOCX / PPTX / XLSX)
# ---------------------------------------------------------------------------
_CORE_TS = re.compile(rb"(<(dcterms:created|dcterms:modified|cp:lastPrinted)\b[^>]*>)[^<]*(</\2>)")


def normalize_zip(path: Path) -> None:
    """Rewrite an OOXML zip with fixed timestamps, entry order and compression."""
    with zipfile.ZipFile(path, "r") as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    fixed: list[tuple[str, bytes]] = []
    for name, data in entries:
        if name == "docProps/core.xml":
            data = _CORE_TS.sub(rb"\g<1>2026-01-01T00:00:00Z\g<3>", data)
        fixed.append((name, data))

    def order(item: tuple[str, bytes]) -> tuple[int, str]:
        name = item[0]
        if name == "[Content_Types].xml":
            return (0, name)
        if name == "_rels/.rels":
            return (1, name)
        return (2, name)

    fixed.sort(key=order)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w") as zout:
        for name, data in fixed:
            info = zipfile.ZipInfo(name, date_time=ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0o600 << 16
            zout.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED, compresslevel=6)
    tmp.replace(path)


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------
def render_pdf(doc: dict, path: Path, resolve_fn) -> None:
    from reportlab import rl_config

    rl_config.invariant = 1
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    base = dict(hyphenationLang="", embeddedHyphenation=0, uriWasteReduce=0, alignment=0)
    st_title = ParagraphStyle(
        "t", fontName="Helvetica-Bold", fontSize=18, leading=22, spaceAfter=6, **base
    )
    st_sub = ParagraphStyle(
        "s", fontName="Helvetica", fontSize=12, leading=15, spaceAfter=4, **base
    )
    st_meta = ParagraphStyle(
        "m", fontName="Helvetica", fontSize=9, leading=12, spaceAfter=2, **base
    )
    st_h1 = ParagraphStyle(
        "h1",
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        spaceBefore=14,
        spaceAfter=8,
        **base,
    )
    st_h2 = ParagraphStyle(
        "h2",
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        spaceBefore=10,
        spaceAfter=6,
        **base,
    )
    st_body = ParagraphStyle(
        "b", fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=9, **base
    )
    st_bullet = ParagraphStyle(
        "bl", fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=4, leftIndent=12, **base
    )
    st_trow = ParagraphStyle(
        "tr", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=2, leftIndent=6, **base
    )
    st_thead = ParagraphStyle(
        "th", fontName="Helvetica-Bold", fontSize=10, leading=13, spaceAfter=2, leftIndent=6, **base
    )

    def para(text: str, style) -> Paragraph:
        return Paragraph(escape(resolve_fn(text)), style)

    story = [para(doc["title"], st_title)]
    if doc.get("subtitle"):
        story.append(para(doc["subtitle"], st_sub))
    for line in doc.get("meta") or []:
        story.append(para(line, st_meta))
    story.append(Paragraph(escape(DISCLAIMER), st_meta))
    story.append(Spacer(1, 10))
    for block in doc["blocks"]:
        kind = block[0]
        if kind == "h1":
            story.append(para(block[1], st_h1))
        elif kind == "h2":
            story.append(para(block[1], st_h2))
        elif kind == "p":
            story.append(para(block[1], st_body))
        elif kind == "bullets":
            for item in block[1]:
                story.append(para("- " + item, st_bullet))
            story.append(Spacer(1, 6))
        elif kind == "table":
            for i, row in enumerate(block[1]):
                story.append(para(" | ".join(row), st_thead if i == 0 else st_trow))
            story.append(Spacer(1, 9))
        elif kind == "pagebreak":
            story.append(PageBreak())
        else:
            raise ValueError(f"unknown block {kind}")

    def footer(canvas, _doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(inch, 0.5 * inch, f"{DISCLAIMER_CORE} - page {canvas.getPageNumber()}")
        canvas.restoreState()

    tpl = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.9 * inch,
        rightMargin=0.9 * inch,
        topMargin=0.9 * inch,
        bottomMargin=0.9 * inch,
        title=doc["title"],
        author=AUTHOR,
        subject=DISCLAIMER,
        creator=AUTHOR,
        producer=AUTHOR,
        invariant=1,
        pageCompression=1,
    )
    tpl.build(story, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------
def render_docx(doc: dict, path: Path, resolve_fn) -> None:
    from docx import Document

    d = Document()
    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.title = doc["title"]
    cp.subject = DISCLAIMER
    cp.created = FIXED_DT
    cp.modified = FIXED_DT
    cp.last_printed = FIXED_DT
    cp.revision = 1
    d.add_paragraph(resolve_fn(doc["title"]), style="Title")
    if doc.get("subtitle"):
        d.add_paragraph(resolve_fn(doc["subtitle"]), style="Normal")
    for line in doc.get("meta") or []:
        d.add_paragraph(resolve_fn(line), style="Normal")
    d.add_paragraph(DISCLAIMER, style="Normal")
    for block in doc["blocks"]:
        kind = block[0]
        if kind == "h1":
            d.add_paragraph(resolve_fn(block[1]), style="Heading 1")
        elif kind == "h2":
            d.add_paragraph(resolve_fn(block[1]), style="Heading 2")
        elif kind == "p":
            d.add_paragraph(resolve_fn(block[1]), style="Normal")
        elif kind == "bullets":
            for item in block[1]:
                d.add_paragraph(resolve_fn(item), style="List Bullet")
        elif kind == "table":
            rows = block[1]
            table = d.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = resolve_fn(cell)
        elif kind == "pagebreak":
            continue
        else:
            raise ValueError(f"unknown block {kind}")
    d.save(str(path))
    normalize_zip(path)


# ---------------------------------------------------------------------------
# PPTX (python-pptx)
# ---------------------------------------------------------------------------
def render_pptx(doc: dict, path: Path, resolve_fn) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    cp = prs.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.title = doc["title"]
    cp.subject = DISCLAIMER
    cp.created = FIXED_DT
    cp.modified = FIXED_DT
    cp.last_printed = FIXED_DT
    cp.revision = 1

    title_layout, content_layout, title_only = (
        prs.slide_layouts[0],
        prs.slide_layouts[1],
        prs.slide_layouts[5],
    )
    s0 = prs.slides.add_slide(title_layout)
    s0.shapes.title.text = resolve_fn(doc["title"])
    sub_tf = s0.placeholders[1].text_frame
    sub_tf.text = resolve_fn(doc.get("subtitle") or "")
    p = sub_tf.add_paragraph()
    p.text = DISCLAIMER

    for spec in doc["slides"]:
        bullets = [resolve_fn(b) for b in spec.get("bullets") or []]
        table = spec.get("table")
        if table:
            slide = prs.slides.add_slide(title_only)
            slide.shapes.title.text = resolve_fn(spec["title"])
            body_h = Inches(0.45 * max(len(bullets), 1) + 0.3) if bullets else Inches(0)
            if bullets:
                box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.1), body_h)
                tf = box.text_frame
                tf.word_wrap = True
                for i, b in enumerate(bullets):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.text = b
                    for run in para.runs:
                        run.font.size = Pt(16)
            n_rows, n_cols = len(table), len(table[0])
            top = Inches(1.6) + body_h
            gshape = slide.shapes.add_table(
                n_rows, n_cols, Inches(0.6), top, Inches(12.1), Inches(0.4 * n_rows)
            )
            for r, row in enumerate(table):
                for c, cell in enumerate(row):
                    tcell = gshape.table.cell(r, c)
                    tcell.text = resolve_fn(cell)
                    for para in tcell.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(12)
        else:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = resolve_fn(spec["title"])
            tf = slide.placeholders[1].text_frame
            for i, b in enumerate(bullets):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = b
                for run in para.runs:
                    run.font.size = Pt(18)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = resolve_fn(spec["notes"])
    prs.save(str(path))
    normalize_zip(path)


# ---------------------------------------------------------------------------
# Markdown / text
# ---------------------------------------------------------------------------
def render_markdown(doc: dict, path: Path, resolve_fn) -> None:
    out: list[str] = [f"# {resolve_fn(doc['title'])}", ""]
    if doc.get("subtitle"):
        out += [resolve_fn(doc["subtitle"]), ""]
    for line in doc.get("meta") or []:
        out += [resolve_fn(line), ""]
    out += [DISCLAIMER, ""]
    for block in doc["blocks"]:
        kind = block[0]
        if kind == "h1":
            out += [f"## {resolve_fn(block[1])}", ""]
        elif kind == "h2":
            out += [f"### {resolve_fn(block[1])}", ""]
        elif kind == "p":
            out += [resolve_fn(block[1]), ""]
        elif kind == "bullets":
            out += [f"- {resolve_fn(item)}" for item in block[1]] + [""]
        elif kind == "table":
            out += [f"- {' | '.join(resolve_fn(c) for c in row)}" for row in block[1]] + [""]
        elif kind == "pagebreak":
            continue
        else:
            raise ValueError(f"unknown block {kind}")
    path.write_text("\n".join(out).rstrip("\n") + "\n", encoding="utf-8", newline="\n")


def render_text(doc: dict, path: Path, resolve_fn) -> None:
    out: list[str] = [resolve_fn(doc["title"]), ""]
    if doc.get("subtitle"):
        out += [resolve_fn(doc["subtitle"]), ""]
    for line in doc.get("meta") or []:
        out += [resolve_fn(line), ""]
    out += [DISCLAIMER, ""]
    for block in doc["blocks"]:
        kind = block[0]
        if kind in ("h1", "h2", "p"):
            out += [resolve_fn(block[1]), ""]
        elif kind == "bullets":
            out += [resolve_fn(item) for item in block[1]] + [""]
        elif kind == "pagebreak":
            continue
        else:
            raise ValueError(f"block {kind} not allowed in text documents")
    path.write_text("\n".join(out).rstrip("\n") + "\n", encoding="utf-8", newline="\n")


# ---------------------------------------------------------------------------
# Tabular
# ---------------------------------------------------------------------------
def render_csv(table: dict, path: Path) -> None:
    encoding = "utf-8-sig" if table.get("bom") else "utf-8"
    with path.open("w", encoding=encoding, newline="") as fh:
        writer = csv.writer(fh, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(table["columns"])
        for row in table["rows"]:
            writer.writerow(["" if v is None else v for v in row])


def render_xlsx(book: dict, path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.title = book["title"]
    wb.properties.subject = DISCLAIMER
    wb.properties.created = FIXED_DT
    wb.properties.modified = FIXED_DT
    for sheet in book["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        ws.append(sheet["columns"])
        for row in sheet["rows"]:
            ws.append(list(row))
    wb.save(str(path))
    normalize_zip(path)
