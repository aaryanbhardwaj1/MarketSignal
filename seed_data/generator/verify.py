"""Self-verification for the MarketSignal seed corpus (fictional data).

Usage (from the repository root):
  uv --directory backend run python ../seed_data/generator/verify.py

Checks:
 (a) manifest files exist, sha256 matches, each < 2 MB, total < 15 MB;
 (b) text extracted with pdfplumber / python-docx / python-pptx (text frames, tables, notes) / openpyxl / csv:
     every sentence anchor appears exactly once per workspace (whitespace-normalised) inside a single extraction
     unit of its own source file; every row anchor matches exactly one row with the expected value; surface forms
     co-occur with the anchor;
 (c) no forbidden real-brand names in generated text, the world model or the generator sources;
 (d) every document contains the fictional disclaimer;
 (e) generate.py --check passes (byte-identical regeneration);
 plus structural checks (ASCII text, BOM, slide/notes/table counts, page counts, ledger quotas).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import subprocess
import sys
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))

from common import (  # noqa: E402
    DISCLAIMER_CORE,
    GEN_DIR,
    SEED_ROOT,
    WORLD_PATH,
    forbidden_hits,
    non_ascii_chars,
    norm_ws,
)

GEN = SEED_ROOT / "generated"
MAX_FILE = 2 * 1024 * 1024
MAX_TOTAL = 15 * 1024 * 1024


class Report:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.notes: list[str] = []

    def fail(self, msg: str) -> None:
        self.errors.append(msg)

    def ok(self, cond: bool, msg: str) -> None:
        if not cond:
            self.fail(msg)


# ---------------------------------------------------------------------------
# Extraction: returns (units, tables, info). units = list of text units.
# tables = {sheet_or_None: (header, rows)}
# ---------------------------------------------------------------------------
def extract_pdf(path: Path) -> tuple[list[str], dict, dict]:
    import pdfplumber

    units = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            units.append(page.extract_text() or "")
        return units, {}, {"pages": len(pdf.pages)}


def extract_docx(path: Path) -> tuple[list[str], dict, dict]:
    from docx import Document

    d = Document(str(path))
    units, styles = [], []
    for p in d.paragraphs:
        units.append(p.text)
        styles.append((p.style.name, p.text))
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                units.append(cell.text)
    return units, {}, {"styles": styles}


def extract_pptx(path: Path) -> tuple[list[str], dict, dict]:
    from pptx import Presentation

    prs = Presentation(str(path))
    units: list[str] = []
    notes = tables = 0
    loc_units: list[tuple[str, str]] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                kind = "title" if shape == slide.shapes.title else "body"
                for para in shape.text_frame.paragraphs:
                    units.append(para.text)
                    loc_units.append((kind, para.text))
            if getattr(shape, "has_table", False) and shape.has_table:
                tables += 1
                for row in shape.table.rows:
                    for cell in row.cells:
                        units.append(cell.text)
                        loc_units.append(("table", cell.text))
        if slide.has_notes_slide:
            txt = slide.notes_slide.notes_text_frame.text
            if txt.strip():
                notes += 1
                units.append(txt)
                loc_units.append(("notes", txt))
    return (
        units,
        {},
        {"slides": len(prs.slides), "notes": notes, "tables": tables, "loc_units": loc_units},
    )


def extract_xlsx(path: Path) -> tuple[list[str], dict, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=False)
    units, tables, merged, typed = [], {}, 0, True
    for ws in wb.worksheets:
        merged += len(ws.merged_cells.ranges)
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        header, body = rows[0], rows[1:]
        tables[ws.title] = (header, body)
        for r in rows:
            units.extend("" if v is None else str(v) for v in r)
        for r in body:
            for v in r:
                if isinstance(v, str) and _looks_numeric(v):
                    typed = False
    return units, tables, {"merged": merged, "typed": typed, "sheets": wb.sheetnames}


def _looks_numeric(v: str) -> bool:
    try:
        float(v)
        return True
    except ValueError:
        return False


def extract_csv(path: Path) -> tuple[list[str], dict, dict]:
    raw = path.read_bytes()
    bom = raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text, newline="")))
    header, body = rows[0], rows[1:]
    units = [c for r in rows for c in r]
    widths = {len(r) for r in rows}
    return units, {None: (header, body)}, {"bom": bom, "widths": widths, "crlf": b"\r\n" in raw}


def extract_text(path: Path) -> tuple[list[str], dict, dict]:
    text = path.read_text(encoding="utf-8")
    return [u for u in text.split("\n\n")], {}, {}


EXTRACT = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "pptx": extract_pptx,
    "xlsx": extract_xlsx,
    "csv": extract_csv,
    "markdown": extract_text,
    "text": extract_text,
}


def _match_value(cell, expected) -> bool:
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        try:
            return abs(float(cell) - float(expected)) < 1e-9
        except (TypeError, ValueError):
            return False
    return str(cell) == str(expected)


def main() -> int:
    rep = Report()
    manifest = json.loads((GEN / "manifest.json").read_text(encoding="utf-8"))
    ledger = json.loads((SEED_ROOT / "fact_ledger.json").read_text(encoding="utf-8"))
    uploads = manifest["uploads"]

    # (a) files, hashes, sizes
    total = 0
    for u in uploads:
        p = GEN / u["filename"]
        if not p.exists():
            rep.fail(f"(a) missing {u['filename']}")
            continue
        size = p.stat().st_size
        total += size
        rep.ok(
            hashlib.sha256(p.read_bytes()).hexdigest() == u["sha256"],
            f"(a) sha256 mismatch {u['filename']}",
        )
        rep.ok(size < MAX_FILE, f"(a) {u['filename']} is {size} bytes (>= 2 MB)")
        import re

        rep.ok(
            re.fullmatch(r"[A-Z0-9]{1,12}(-[A-Z0-9]{1,12}){0,5}", u["source_code"]) is not None,
            f"(a) bad source_code {u['source_code']}",
        )
    rep.ok(total < MAX_TOTAL, f"(a) total size {total} >= 15 MB")
    on_disk = sorted(
        str(f.relative_to(GEN)) for f in GEN.rglob("*") if f.is_file() and f.name != "manifest.json"
    )
    rep.ok(
        on_disk == sorted(u["filename"] for u in uploads),
        "(a) generated/ contains files not in manifest",
    )

    # extract everything
    extracted: dict[str, tuple] = {}
    for u in uploads:
        extracted[u["filename"]] = EXTRACT[u["source_type"]](GEN / u["filename"])

    # (c) forbidden names, (d) disclaimer, ASCII
    for u in uploads:
        units, _, info = extracted[u["filename"]]
        text = "\n".join(units)
        hits = forbidden_hits(text)
        rep.ok(not hits, f"(c) forbidden terms {hits} in {u['filename']}")
        rep.ok(DISCLAIMER_CORE in norm_ws(text), f"(d) disclaimer missing in {u['filename']}")
        bad = non_ascii_chars(text)
        rep.ok(not bad, f"ASCII: non-ASCII {bad!r} in {u['filename']}")
    for extra in [
        WORLD_PATH,
        *sorted(GEN_DIR.glob("*.py")),
        *sorted((GEN_DIR / "content").glob("*.py")),
        SEED_ROOT / "fact_ledger.json",
        GEN / "manifest.json",
        SEED_ROOT / "README.md",
    ]:
        if extra.name == "common.py":
            continue  # holds the forbidden-term list itself
        hits = forbidden_hits(extra.read_text(encoding="utf-8"))
        rep.ok(not hits, f"(c) forbidden terms {hits} in {extra.relative_to(SEED_ROOT)}")

    # (b) anchors
    by_ws_units: dict[str, list[tuple[str, str]]] = {}
    for u in uploads:
        units = extracted[u["filename"]][0]
        by_ws_units.setdefault(u["workspace_code"], []).extend(
            (u["filename"], norm_ws(x)) for x in units
        )
    for f in ledger["facts"]:
        a = f["anchor"]
        if a["type"] == "sentence":
            needle = norm_ws(a["text"])
            hits = [
                (fn, unit.count(needle))
                for fn, unit in by_ws_units[f["workspace_code"]]
                if needle in unit
            ]
            n = sum(c for _, c in hits)
            rep.ok(
                n == 1,
                f"(b) {f['fact_id']} sentence anchor found {n} times in {f['workspace_code']}: {hits[:3]}",
            )
            if n == 1:
                rep.ok(
                    hits[0][0] == f["filename"],
                    f"(b) {f['fact_id']} anchor in {hits[0][0]}, expected {f['filename']}",
                )
            for sf in f["surface_forms"]:
                rep.ok(
                    sf.lower() in needle.lower(),
                    f"(b) {f['fact_id']} surface form {sf!r} not in anchor",
                )
            if f["source_type"] == "pptx" and f.get("location") and n == 1:
                locs = extracted[f["filename"]][2]["loc_units"]
                ok_loc = any(kind == f["location"] and needle in norm_ws(t) for kind, t in locs)
                rep.ok(ok_loc, f"(b) {f['fact_id']} not in pptx {f['location']}")
            if f["source_code"] == "INTERVIEWS" and n == 1:
                styles = extracted[f["filename"]][2]["styles"]
                rep.ok(
                    any(t.startswith("A: ") and needle in norm_ws(t) for _, t in styles),
                    f"(b) {f['fact_id']} not inside an 'A: ' answer",
                )
        else:
            tables = extracted[f["filename"]][1]
            if a["sheet"] not in tables:
                rep.fail(f"(b) {f['fact_id']} sheet {a['sheet']} missing")
                continue
            header, body = tables[a["sheet"]]
            try:
                ki, vi = header.index(a["key_column"]), header.index(a["value_column"])
            except ValueError:
                rep.fail(f"(b) {f['fact_id']} columns missing")
                continue
            rows = [r for r in body if str(r[ki]) == str(a["key_value"])]
            rep.ok(len(rows) == 1, f"(b) {f['fact_id']} row key matched {len(rows)} rows")
            if len(rows) == 1:
                rep.ok(
                    _match_value(rows[0][vi], a["value"]),
                    f"(b) {f['fact_id']} value {rows[0][vi]!r} != {a['value']!r}",
                )
                joined = " ".join("" if v is None else str(v) for v in rows[0])
                for sf in f["surface_forms"]:
                    rep.ok(
                        sf.lower() in joined.lower(),
                        f"(b) {f['fact_id']} surface form {sf!r} not in row",
                    )

    # canary isolation: Southpeak canary values must not appear in Northstar text (anchors / injection codes)
    ns_text = " ".join(t for _, t in by_ws_units["NORTHSTAR"])
    for f in ledger["facts"]:
        if f["is_canary"] and f["anchor"]["type"] == "sentence":
            rep.ok(
                norm_ws(f["anchor"]["text"]) not in ns_text,
                f"canary {f['fact_id']} leaked into NORTHSTAR",
            )
    rep.ok("Southpeak" not in ns_text, "Southpeak mentioned in NORTHSTAR corpus")
    rep.ok(
        "Northstar" not in " ".join(t for _, t in by_ws_units["SOUTHPEAK"]),
        "Northstar mentioned in SOUTHPEAK",
    )

    # structural checks
    for u in uploads:
        units, tables, info = extracted[u["filename"]]
        fn, st = u["filename"], u["source_type"]
        if st == "pdf":
            lo, hi = (
                (6, 8)
                if u["source_code"] == "BRAND-STRATEGY" and u["workspace_code"] == "NORTHSTAR"
                else (3, 8)
            )
            rep.ok(lo <= info["pages"] <= hi, f"{fn}: {info['pages']} pages, expected {lo}-{hi}")
            rep.notes.append(f"{fn}: {info['pages']} pages")
        if st == "pptx":
            rep.notes.append(
                f"{fn}: {info['slides']} slides, {info['notes']} with notes, {info['tables']} tables"
            )
            rep.ok(
                info["notes"] >= 4 and info["tables"] >= 1, f"{fn}: needs >=4 notes and >=1 table"
            )
            if u["source_code"] == "Q3-REVIEW":
                rep.ok(12 <= info["slides"] <= 15, f"{fn}: {info['slides']} slides, expected 12-15")
        if st == "xlsx":
            rep.ok(info["merged"] == 0, f"{fn}: merged cells present")
            rep.ok(info["typed"], f"{fn}: numeric-looking strings found (numbers must be typed)")
            for name, (_, body) in tables.items():
                rep.notes.append(f"{fn}[{name}]: {len(body)} rows")
        if st == "csv":
            rep.ok(len(info["widths"]) == 1, f"{fn}: ragged rows")
            rep.ok(info["crlf"], f"{fn}: expected CRLF line endings")
            rep.notes.append(f"{fn}: {len(tables[None][1])} rows")
            if u["source_code"] == "SURVEY-2026" and u["workspace_code"] == "NORTHSTAR":
                rep.ok(info["bom"], f"{fn}: survey must have a UTF-8 BOM")
    expected_sheets = {
        "northstar/Northstar_Product_Performance.xlsx": [
            "Category_Monthly",
            "SKU_Performance",
            "Returns",
        ]
    }
    for fn, sheets in expected_sheets.items():
        rep.ok(extracted[fn][2]["sheets"] == sheets, f"{fn}: sheets {extracted[fn][2]['sheets']}")
    styles = extracted["northstar/Northstar_Customer_Interviews.docx"][2]["styles"]
    n_int = sum(1 for s, t in styles if s == "Heading 1" and t.startswith("Interview "))
    rep.ok(n_int == 12, f"interviews docx has {n_int} interview sections")

    # ledger quotas
    c = ledger["counts"]
    by_t = c["northstar_core_by_source_type"]
    quotas = {"pdf": 10, "docx": 8, "pptx": 8, "xlsx": 8, "csv": 8, "markdown": 6, "text": 2}
    for t, q in quotas.items():
        rep.ok(by_t.get(t, 0) >= q, f"ledger: {t} facts {by_t.get(t, 0)} < {q}")
    core = [
        f
        for f in ledger["facts"]
        if f["workspace_code"] == "NORTHSTAR"
        and not f["distractor_of"]
        and f["category"] != "injection"
    ]
    rep.ok(len(core) >= 60, "ledger: < 60 Northstar facts")
    rep.ok(
        sum(1 for f in core if f["source_code"] == "INTERVIEWS") >= 5, "ledger: < 5 interview facts"
    )
    pp = [f for f in core if f["source_type"] == "pptx"]
    rep.ok(sum(1 for f in pp if f["location"] == "notes") >= 3, "ledger: < 3 pptx notes facts")
    rep.ok(sum(1 for f in pp if f["location"] == "table") >= 2, "ledger: < 2 pptx table facts")
    rep.ok(
        c["canaries"] >= 6
        and c["injections"] == 2
        and c["contradiction_pairs"] >= 4
        and c["superseded"] >= 3
        and c["distractors"] >= 8,
        f"ledger counts too low: {c}",
    )
    ids = {f["fact_id"] for f in ledger["facts"]}
    for pair in ledger["contradictions"]:
        rep.ok(
            pair["a"] in ids and pair["b"] in ids,
            f"contradiction {pair['id']} refers to unknown facts",
        )

    # (e) byte-identical regeneration
    proc = subprocess.run(
        [sys.executable, str(GEN_DIR / "generate.py"), "--check"], capture_output=True, text=True
    )
    rep.ok(
        proc.returncode == 0,
        f"(e) generate.py --check failed: {proc.stdout[-2000:]} {proc.stderr[-2000:]}",
    )

    for n in rep.notes:
        print("  ", n)
    for e in rep.errors:
        print("FAIL", e)
    print(f"verify: {len(uploads)} uploads, {total} bytes, {len(ledger['facts'])} ledger facts")
    print("verify OK" if not rep.errors else f"verify FAILED ({len(rep.errors)} problems)")
    return 1 if rep.errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
